from datetime import datetime
from pathlib import Path
from tkinter import filedialog, messagebox, ttk

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph


def _filtered_rows(tree, criterion, start_date, end_date):
    columns = list(tree["columns"])
    rows = []
    for item_id in tree.get_children():
        values = tuple(tree.item(item_id, "values"))
        text = " ".join(str(value) for value in values).lower()
        if criterion and criterion.lower() not in text:
            continue
        date_values = [str(values[index])[:10] for index, column in enumerate(columns)
                       if "fecha" in column.lower() or "date" in column.lower()]
        if start_date or end_date:
            if not date_values:
                continue
            if start_date and all(value < start_date for value in date_values):
                continue
            if end_date and all(value > end_date for value in date_values):
                continue
        rows.append(values)
    return columns, rows


def _ask_filter(parent):
    dialog = ttk.Frame(parent)
    window = dialog.winfo_toplevel()
    filter_window = ttk.Frame(window)
    popup = __import__("tkinter").Toplevel(window)
    popup.title("Filtros de exportación")
    popup.transient(window)
    popup.grab_set()
    ttk.Label(popup, text="Criterio (opcional):").grid(row=0, column=0, padx=8, pady=6, sticky="w")
    criterion = ttk.Entry(popup, width=28)
    criterion.grid(row=0, column=1, padx=8, pady=6)
    ttk.Label(popup, text="Desde (AAAA-MM-DD):").grid(row=1, column=0, padx=8, pady=6, sticky="w")
    start = ttk.Entry(popup, width=28)
    start.grid(row=1, column=1, padx=8, pady=6)
    ttk.Label(popup, text="Hasta (AAAA-MM-DD):").grid(row=2, column=0, padx=8, pady=6, sticky="w")
    end = ttk.Entry(popup, width=28)
    end.grid(row=2, column=1, padx=8, pady=6)
    result = {}

    def accept():
        for label, entry in (("Desde", start), ("Hasta", end)):
            value = entry.get().strip()
            if value:
                try:
                    datetime.strptime(value, "%Y-%m-%d")
                except ValueError:
                    messagebox.showerror("Filtro inválido", f"{label}: usa AAAA-MM-DD.", parent=popup)
                    return
        result.update(criterion=criterion.get().strip(), start=start.get().strip(), end=end.get().strip())
        popup.destroy()

    ttk.Button(popup, text="Aplicar", command=accept, style="Accion.TButton").grid(row=3, column=0, columnspan=2, pady=10)
    popup.wait_window()
    return result if result else None


def _get_export_data(parent, tree):
    filters = _ask_filter(parent)
    if filters is None:
        return None
    return _filtered_rows(tree, filters["criterion"], filters["start"], filters["end"])


def export_tree_to_excel(parent, tree, title):
    data = _get_export_data(parent, tree)
    if data is None:
        return
    columns, rows = data
    path = filedialog.asksaveasfilename(parent=parent, defaultextension=".xlsx", filetypes=[("Excel", "*.xlsx")], initialfile=f"{title}.xlsx")
    if not path:
        return
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = title[:31]
    sheet.append(columns)
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="007ACC")
    for row in rows:
        sheet.append(row)
    for column in sheet.columns:
        width = min(max(len(str(cell.value or "")) for cell in column) + 2, 40)
        sheet.column_dimensions[column[0].column_letter].width = width
    workbook.save(path)
    messagebox.showinfo("Exportación", f"Archivo Excel creado:\n{Path(path).name}", parent=parent)


def export_tree_to_pdf(parent, tree, title):
    data = _get_export_data(parent, tree)
    if data is None:
        return
    columns, rows = data
    path = filedialog.asksaveasfilename(parent=parent, defaultextension=".pdf", filetypes=[("PDF", "*.pdf")], initialfile=f"{title}.pdf")
    if not path:
        return
    document = SimpleDocTemplate(path, pagesize=landscape(A4), rightMargin=24, leftMargin=24, topMargin=24, bottomMargin=24)
    styles = getSampleStyleSheet()
    heading = Paragraph(title, styles["Title"])
    table = Table([columns, *rows], repeatRows=1)
    table.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#007ACC")),
        ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
        ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
        ("GRID", (0, 0), (-1, -1), 0.25, colors.grey),
        ("FONTSIZE", (0, 0), (-1, -1), 7),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EAF3F8")]),
    ]))
    document.build([heading, table])
    messagebox.showinfo("Exportación", f"Archivo PDF creado:\n{Path(path).name}", parent=parent)


def add_export_controls(parent, tree, title):
    frame = ttk.Frame(parent)
    frame.pack(fill="x", padx=8, pady=(4, 6), before=tree.master)
    ttk.Label(frame, text="Exportar filtrado:", style="Campo.TLabel").pack(side="left", padx=(0, 8))
    ttk.Button(frame, text="📊 Excel", command=lambda: export_tree_to_excel(parent, tree, title), style="Accion.TButton").pack(side="left", padx=3)
    ttk.Button(frame, text="📄 PDF", command=lambda: export_tree_to_pdf(parent, tree, title), style="Accion.TButton").pack(side="left", padx=3)
